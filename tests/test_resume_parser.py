"""Tests for resume_parser.extract_text — PDF / DOCX / TXT extraction."""

import io

import pytest
from docx import Document
from fpdf import FPDF

from resume_parser import extract_text


class FakeUploadedFile:
    """Minimal stand-in for Streamlit's UploadedFile (.name / .getvalue())."""

    def __init__(self, name: str, data: bytes):
        self.name = name
        self._data = data

    def getvalue(self) -> bytes:
        return self._data


# --------------------------------------------------------------------------
# TXT
# --------------------------------------------------------------------------


def test_txt_extraction_returns_raw_text():
    upload = FakeUploadedFile("resume.txt", b"Jane Doe - Software Engineer")
    assert extract_text(upload) == "Jane Doe - Software Engineer"


def test_txt_extraction_tolerates_bad_bytes():
    # A byte sequence that isn't valid UTF-8 shouldn't crash extraction.
    upload = FakeUploadedFile("resume.txt", b"Valid text \xff\xfe more text")
    result = extract_text(upload)
    assert "Valid text" in result
    assert "more text" in result


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------


def _build_docx_bytes(paragraphs=None, table_rows=None) -> bytes:
    doc = Document()
    for text in paragraphs or []:
        doc.add_paragraph(text)
    if table_rows:
        table = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, cell_text in enumerate(row):
                table.rows[r].cells[c].text = cell_text
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_extraction_pulls_paragraph_text():
    data = _build_docx_bytes(paragraphs=["Jane Doe", "Software Engineer"])
    upload = FakeUploadedFile("resume.docx", data)
    text = extract_text(upload)
    assert "Jane Doe" in text
    assert "Software Engineer" in text


def test_docx_extraction_pulls_table_cell_text():
    # Resumes often use tables for layout (e.g. a two-column skills section) —
    # extract_text must not silently skip that content.
    data = _build_docx_bytes(
        paragraphs=["Jane Doe"],
        table_rows=[["Skills", "Python, SQL, Docker"]],
    )
    upload = FakeUploadedFile("resume.docx", data)
    text = extract_text(upload)
    assert "Skills" in text
    assert "Python, SQL, Docker" in text


def test_docx_with_no_text_raises_value_error():
    data = _build_docx_bytes(paragraphs=[])
    upload = FakeUploadedFile("resume.docx", data)
    with pytest.raises(ValueError, match="No extractable text"):
        extract_text(upload)


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------


def _build_pdf_bytes(text: str = None) -> bytes:
    pdf = FPDF()
    pdf.add_page()
    if text:
        pdf.set_font("Helvetica", size=12)
        pdf.cell(0, 10, text)
    return bytes(pdf.output())


def test_pdf_extraction_returns_text():
    data = _build_pdf_bytes("Jane Doe - PDF Resume")
    upload = FakeUploadedFile("resume.pdf", data)
    assert "Jane Doe - PDF Resume" in extract_text(upload)


def test_pdf_with_no_text_raises_value_error():
    # A page with no text content at all (e.g. a scanned-image PDF) should
    # give an actionable error rather than silently returning an empty string.
    data = _build_pdf_bytes(text=None)
    upload = FakeUploadedFile("resume.pdf", data)
    with pytest.raises(ValueError, match="scanned image"):
        extract_text(upload)


# --------------------------------------------------------------------------
# Unsupported / edge cases
# --------------------------------------------------------------------------


def test_unsupported_extension_raises_value_error():
    upload = FakeUploadedFile("resume.xyz", b"whatever")
    with pytest.raises(ValueError, match="Unsupported file type"):
        extract_text(upload)


def test_extension_matching_is_case_insensitive():
    upload = FakeUploadedFile("RESUME.TXT", b"Case insensitive check")
    assert extract_text(upload) == "Case insensitive check"
